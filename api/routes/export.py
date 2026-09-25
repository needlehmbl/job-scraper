"""Offline spreadsheet export for the dashboard's Download button.

GET /export/xlsx -- Jobs + Filtered tabs as a styled .xlsx workbook.

Styling echoes the dashboard: dark header row, frozen panes, autofilter,
status tints matching the dashboard pills, and clickable job links.
"""
import io
from datetime import date, datetime

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from api import db

router = APIRouter()

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="2563EB", underline="single")
WRAP = Alignment(vertical="center", wrap_text=True)

# Dashboard STATUS_STYLES, translated to fills.
STATUS_FILLS = {
    "NEW": PatternFill("solid", fgColor="E5E7EB"),
    "REVIEWED": PatternFill("solid", fgColor="FEF3C7"),
    "APPLIED": PatternFill("solid", fgColor="D1FAE5"),
    "SKIP": PatternFill("solid", fgColor="F3F4F6"),
    "MISMATCH": PatternFill("solid", fgColor="FFEDD5"),
    "EXP_GAP": PatternFill("solid", fgColor="EDE9FE"),
    "EXPIRED": PatternFill("solid", fgColor="D6D3D1"),
    "DUPLICATE": PatternFill("solid", fgColor="F3F4F6"),
}

JOBS_COLUMNS = [
    ("title", "Title", 45),
    ("company", "Company", 28),
    ("score", "Fit", 8),
    ("status", "Status", 12),
    ("stage", "Stage", 12),
    ("source", "Source", 12),
    ("location", "Location", 26),
    ("date_posted", "Posted", 12),
    ("scraped_at", "Scraped", 17),
    ("follow_up_at", "Follow-up", 12),
    ("url", "Link", 18),
]

FILTERED_COLUMNS = [
    ("title", "Title", 45),
    ("company", "Company", 28),
    ("source", "Source", 12),
    ("filter_reason", "Reason", 30),
    ("search_term", "Search term", 24),
    ("location", "Location", 26),
    ("url", "Link", 18),
    ("filtered_at", "Filtered at", 17),
]


def _text(value) -> str | int | float:
    """Excel-safe scalar: dates become ISO strings, None becomes ''."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return value
    return str(value)


def _write_sheet(ws, columns, rows: list[dict]):
    for ci, (_, header, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=ci, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = width
    for ri, row in enumerate(rows, start=2):
        for ci, (key, _, _) in enumerate(columns, start=1):
            value = row.get(key)
            cell = ws.cell(row=ri, column=ci, value=_text(value))
            cell.alignment = WRAP
            if key == "status" and value in STATUS_FILLS:
                cell.fill = STATUS_FILLS[value]
            if key == "url" and value:
                cell.hyperlink = str(value)
                cell.value = "open"
                cell.font = LINK_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


@router.get("/export/xlsx")
def export_xlsx():
    with db.connect() as conn, conn.cursor() as cur:
        cur.execute(
            """
            SELECT id, title, company, score, status, stage, source,
                   location, date_posted, scraped_at, follow_up_at, url
            FROM jobs
            ORDER BY score DESC, scraped_at DESC
            """
        )
        cols = [d.name for d in cur.description]
        jobs = [dict(zip(cols, r)) for r in cur.fetchall()]
    try:
        filtered = db.list_filtered_jobs()
    except Exception:
        filtered = []

    wb = Workbook()
    ws_jobs = wb.active
    ws_jobs.title = "Jobs"
    _write_sheet(ws_jobs, JOBS_COLUMNS, jobs)
    ws_filt = wb.create_sheet("Filtered")
    _write_sheet(ws_filt, FILTERED_COLUMNS, filtered)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y-%m-%d")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="jobs-{stamp}.xlsx"'},
    )
